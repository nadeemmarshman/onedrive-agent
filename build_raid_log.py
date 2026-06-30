from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
BODY_FONT = Font(name=FONT_NAME, size=10)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
THIN_BORDER = Border(*[Side(style="thin", color="D9D9D9")] * 4)

STATUS_COLORS = {
    "Open": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
    "Monitoring": PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7"),
    "Closed": PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA"),
    "Active": PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7"),
    "Pending": PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
}


def style_header_row(ws, row_num, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        cell.border = THIN_BORDER


def write_sheet(ws, title, subtitle, headers, rows, col_widths, status_col_index=None):
    ws.sheet_view.showGridLines = False

    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))

    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)
    style_header_row(ws, header_row, len(headers))

    for r_offset, row_data in enumerate(rows, start=1):
        row_num = header_row + r_offset
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if status_col_index is not None and col_idx == status_col_index:
                status_key = next((k for k in STATUS_COLORS if value.startswith(k)), None)
                if status_key:
                    cell.fill = STATUS_COLORS[status_key]
                    cell.font = Font(name=FONT_NAME, size=10, bold=True)

    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"


wb = Workbook()

# --- Overview sheet ---
ws_overview = wb.active
ws_overview.title = "Overview"
ws_overview.sheet_view.showGridLines = False
ws_overview["A1"] = "OneDrive Agent — RAID Log"
ws_overview["A1"].font = TITLE_FONT
ws_overview["A2"] = "Risks, Assumptions, Issues, and Dependencies tracked for this project."
ws_overview["A2"].font = SUBTITLE_FONT
ws_overview["A4"] = (
    "Separate from BACKLOG.md (Agile/Scrum delivery tracking) and the handoff "
    "document (agreed project state and standing rules). RAID tracks what could "
    "go wrong, what we're assuming, what has gone wrong, and what we depend on."
)
ws_overview["A4"].font = BODY_FONT
ws_overview["A4"].alignment = Alignment(wrap_text=True, vertical="top")
ws_overview.merge_cells("A4:F4")
ws_overview.row_dimensions[4].height = 45

summary_data = [
    ("Risks", 4, "See Risks tab"),
    ("Assumptions", 4, "See Assumptions tab"),
    ("Issues", 3, "See Issues tab"),
    ("Dependencies", 4, "See Dependencies tab"),
]
ws_overview["A6"] = "Category"
ws_overview["B6"] = "Items logged"
ws_overview["C6"] = "Detail"
style_header_row(ws_overview, 6, 3)
for i, (cat, count, detail) in enumerate(summary_data, start=7):
    ws_overview.cell(row=i, column=1, value=cat).font = BODY_FONT
    ws_overview.cell(row=i, column=2, value=count).font = BODY_FONT
    ws_overview.cell(row=i, column=3, value=detail).font = BODY_FONT
    for c in range(1, 4):
        ws_overview.cell(row=i, column=c).border = THIN_BORDER

ws_overview.column_dimensions["A"].width = 20
ws_overview.column_dimensions["B"].width = 15
ws_overview.column_dimensions["C"].width = 25

# --- Risks sheet ---
risks_headers = ["#", "Raised", "Risk", "Likelihood", "Impact", "Mitigation", "Status"]
risks_rows = [
    ["R1", "2026-06-28",
     "Load shedding / power loss interrupts the agent mid-loop, losing progress or leaving an action in an unsafe partial state.",
     "High (recurring, locale-specific operational reality in South Africa)",
     "Medium-High — could lose work, or worse, resume incorrectly into a destructive action without approval.",
     "Phase 4.5: explicit state-machine design with persisted state; AWAITING_HUMAN_APPROVAL must never be silently skipped on resume.",
     "Open — design decided (handoff doc v4.0), not yet implemented"],
    ["R2", "2026-06-30",
     "Verbal/status-only confirmation between sessions could drift from actual file/repo/Project-knowledge state, since there is no automatic way to verify \"done\" claims.",
     "Medium (already materialized once — see Issue I1)",
     "Medium — could lead to working from a stale or incorrect document/codebase without realizing it.",
     "Screenshot-verification standing rule adopted (handoff doc v4.0) — checklist items require visual evidence, not verbal confirmation alone.",
     "Monitoring — mitigation in place, watching for recurrence"],
    ["R3", "2026-06-28",
     "Real API/network failures (rate limits, connectivity, insufficient credit) during a live Anthropic API call.",
     "Medium (normal for any real API integration)",
     "Low-Medium — would crash the script ungracefully if unhandled.",
     "Specific exception handling implemented in decision_loop.py (Phase 3, v3.2) for auth, rate limit, connection, and general API errors.",
     "Closed — mitigated and verified (simulated test of each branch)"],
    ["R4", "2026-06-28",
     "An LLM-driven decision loop could return zero, multiple, or unrecognized tool calls, or malformed arguments, none of which a \"happy path only\" execution loop would handle correctly.",
     "High (normal LLM agent behavior, already observed once — Claude returned two tool calls in the same turn during Phase 3 testing)",
     "Medium — could silently skip work, crash, or call an undefined function.",
     "Named explicitly as a proactive-design standing rule (handoff doc v3.1); to be implemented in Phase 4's execution loop.",
     "Open — design decided, implementation pending (Phase 4)"],
]
ws_risks = wb.create_sheet("Risks")
write_sheet(ws_risks, "Risks", "Things that might happen, with a negative impact if they do.",
            risks_headers, risks_rows,
            col_widths=[6, 12, 45, 25, 35, 45, 30], status_col_index=7)

# --- Assumptions sheet ---
assumptions_headers = ["#", "Raised", "Assumption", "Risk if wrong", "Status"]
assumptions_rows = [
    ["A1", "2026-06-30",
     "A successful git commit/push (no error output) means the intended, correct file content was committed.",
     "Proven false — see Issue I1. A commit can succeed with completely wrong content (e.g. an accidental overwrite) and show no error at all.",
     "Closed — disproven; replaced by checking commit insertion/deletion stats and verifying actual committed content (e.g. git show HEAD:file)."],
    ["A2", "2026-06-26",
     "SendGrid's free tier (100 emails/day) will be sufficient for this project's alerting volume.",
     "Low risk if wrong — alert volume for a personal portfolio project should be very low; would only become a problem under a failure loop sending repeated alerts.",
     "Monitoring"],
    ["A3", "2026-06-26",
     "The $6 / minimum Anthropic API credit balance will be sufficient to complete Phases 3 through 7 of testing and development.",
     "Low-Medium — if wrong, would require topping up credit; not a functional blocker, just a cost one.",
     "Monitoring"],
    ["A4", "2026-06-26",
     "Python 3.14 (a very new release) will continue to have adequate third-party package support (prebuilt wheels) for the libraries this project needs.",
     "Medium if wrong — already caused one build failure (tokenizers/Rust compiler issue), resolved by updating pip; could recur with a different package.",
     "Monitoring"],
]
ws_assumptions = wb.create_sheet("Assumptions")
write_sheet(ws_assumptions, "Assumptions", "Things taken as true without proof, which could turn out to be wrong.",
            assumptions_headers, assumptions_rows,
            col_widths=[6, 12, 45, 50, 40], status_col_index=5)

# --- Issues sheet ---
issues_headers = ["#", "Raised", "Issue", "Root cause", "Resolution", "Status"]
issues_rows = [
    ["I1", "2026-06-30",
     "README.md and BACKLOG.md were overwritten with one-line placeholder text (via echo \"...\" > file) instead of containing their real, previously-generated content, and this was committed and pushed to GitHub before being caught.",
     "Combination of Risk R2 (status drift between sessions) and Assumption A1 (wrongly trusting a clean commit/push as proof of correct content) — the actual files on disk were placeholder versions at the time of commit.",
     "Caught via the newly-adopted screenshot-verification rule, by noticing the commit's 84 deletions(-) stat didn't match what should have been a pure-addition commit. Real files re-downloaded, re-saved, re-committed (198 insertions(+), 2 deletions(-) — correct signature), and verified directly via git show HEAD:README.md.",
     "Closed — resolved 2026-06-30, verified by content inspection, not just commit success"],
    ["I2", "2026-06-26",
     "pip install anthropic failed during Phase 3 setup: the tokenizers dependency attempted to compile from Rust source, and no Rust compiler was present on the machine.",
     "Python 3.14 is a very new release; a prebuilt wheel for tokenizers may not yet have been available for it at the time, causing pip to fall back to a source build.",
     "Updated pip first (python -m pip install --upgrade pip), after which a prebuilt wheel installed cleanly without needing Rust.",
     "Closed — resolved 2026-06-26"],
    ["I3", "2026-06-30",
     "Confusion over which handoff document version (v3.0 vs v3.2) was actually loaded in Project knowledge, due to an intermediate swap that wasn't explicitly confirmed back in chat.",
     "Same root cause as R2/A1 — reliance on verbal/remembered status rather than verified evidence.",
     "Clarified via direct question and screenshot; correct version (v4.0) confirmed loaded via Project knowledge files panel screenshot.",
     "Closed — resolved 2026-06-29/30"],
]
ws_issues = wb.create_sheet("Issues")
write_sheet(ws_issues, "Issues", "Things that have already happened and required resolution.",
            issues_headers, issues_rows,
            col_widths=[6, 12, 55, 45, 55, 35], status_col_index=6)

# --- Dependencies sheet ---
deps_headers = ["#", "Raised", "Dependency", "Why it matters", "Status"]
deps_rows = [
    ["D1", "2026-06-26", "Anthropic API availability and billing (console.anthropic.com)",
     "Phase 3 onward cannot function without a live, funded API key.", "Active — funded, key created and working"],
    ["D2", "2026-06-26", "GitHub availability, as the portfolio repo's host",
     "The actual shareable deliverable lives here; repo visibility (Phase 6) also depends on GitHub.", "Active"],
    ["D3", "2026-06-28", "SendGrid account, sender verification, and API availability",
     "Phase 4.5's email alerting cannot function without this.", "Pending — account setup not yet complete"],
    ["D4", "2026-06-26", "Local machine stability / power supply",
     "The agent's own resilience design (Phase 4.5) exists specifically because this dependency is unreliable (load shedding) — see Risk R1.",
     "Active — ongoing operational reality, not something to resolve, only design around"],
]
ws_deps = wb.create_sheet("Dependencies")
write_sheet(ws_deps, "Dependencies", "Things this project relies on, often external, that aren't directly controlled.",
            deps_headers, deps_rows,
            col_widths=[6, 12, 40, 55, 40], status_col_index=5)

wb.save("/home/claude/agent_project/RAID_Log.xlsx")
print("Saved RAID_Log.xlsx")
